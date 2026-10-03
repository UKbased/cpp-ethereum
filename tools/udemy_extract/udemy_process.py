#!/usr/bin/env python3
"""udemy_process.py - turn a Udemy course manifest into a local, searchable knowledge base.

Personal-use companion to udemy_manifest.js and udemy_frames.js. Everything it fetches is
stored on this machine only. It downloads caption files, resources and slide images from the
signed links in the manifest (no login cookies are used or needed), extracts their text,
OCRs frames and images, and writes per-lecture markdown plus one course.md.

Subcommands (run with -h for options):
  build MANIFEST --out DIR       captions, article text, slides, resources, quizzes -> DIR
  sprite-times DIR               slide-change timestamps from Udemy thumbnail sprites
  frame-plan DIR                 video lectures still lacking frames, with URL and udfStart(...) call
  frames DIR FRAMES_JSON...      store + OCR frames saved by udemy_frames.js, then reassemble
  skip-frames DIR --lecture ID   record that frames cannot be captured for a lecture (e.g. DRM)
  screen-capture DIR --lecture ID --seconds N   fallback: grab the screen while a lecture plays
  transcribe DIR --lecture ID (--media FILE | --download)   Whisper for lectures without captions
  assemble DIR                   rebuild course.md and lectures/*.md from the saved state
  status DIR                     what is done and what is still missing
"""
from __future__ import annotations

import argparse
import base64
import datetime as dt
import difflib
import glob
import html
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from pathlib import Path

VERSION = "1.0.0"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
MEDIA_EXT = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm", ".wmv", ".flv",
             ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".wma"}
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg",
            ".py", ".r", ".js", ".ts", ".sql", ".ipynb", ".pine", ".mq4", ".mq5", ".mqh",
            ".afl", ".efs", ".cs", ".vb", ".bas", ".java", ".cpp", ".c", ".h", ".rtf", ".log"}
EMBED_LIMIT = 15000      # chars of each resource's text embedded in course.md
WIN_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)),
                *(f"lpt{i}" for i in range(1, 10))}

try:  # never crash on a narrow Windows console code page
    sys.stdout.reconfigure(errors="replace")
    sys.stderr.reconfigure(errors="replace")
except Exception:
    pass


# ----------------------------------------------------------------------------- helpers
def log(msg: str) -> None:
    print(msg, flush=True)


def read_json(p) -> object:
    with open(p, "r", encoding="utf-8") as f:
        return json.load(f)


def write_json(p, obj) -> None:
    p = Path(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_name(p.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
    os.replace(tmp, p)


def read_text(p, limit: int | None = None) -> str:
    data = Path(p).read_bytes()
    if limit:
        data = data[:limit]
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def safe_name(s: str, maxlen: int = 60) -> str:
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", str(s or "")).strip()
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"_+", "_", s).strip("._ ")
    if not s:
        s = "untitled"
    if s.split(".")[0].lower() in WIN_RESERVED:
        s = "_" + s
    if len(s) > maxlen:
        stem, ext = os.path.splitext(s)
        ext = ext if len(ext) <= 8 else ""
        s = stem[: maxlen - len(ext)].rstrip("._ ") + ext
    return s


def fmt_ts(sec) -> str:
    sec = int(float(sec or 0))
    h, rem = divmod(sec, 3600)
    m, s = divmod(rem, 60)
    return f"{h:d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def rel(p: Path, base: Path) -> str:
    try:
        return Path(p).relative_to(base).as_posix()
    except ValueError:
        return str(p)


def similar(a: str, b: str) -> float:
    a, b = " ".join(a.split()), " ".join(b.split())
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a[:4000], b[:4000]).ratio()


def download(url: str, dest: Path, timeout: int = 120, retries: int = 3,
             max_bytes: int = 2_000_000_000) -> tuple[bool, str | None]:
    """Download url to dest (skipped if dest already exists). Returns (ok, error)."""
    dest = Path(dest)
    if dest.exists() and dest.stat().st_size > 0:
        return True, None
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_name(dest.name + ".part")
    last = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Referer": "https://www.udemy.com/"})
            with urllib.request.urlopen(req, timeout=timeout) as r, open(tmp, "wb") as f:
                n = 0
                while True:
                    chunk = r.read(1 << 16)
                    if not chunk:
                        break
                    n += len(chunk)
                    if n > max_bytes:
                        raise IOError(f"file larger than {max_bytes} bytes")
                    f.write(chunk)
            os.replace(tmp, dest)
            time.sleep(0.2)
            return True, None
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}"
            if e.code in (401, 403, 410):
                last += " (signed link probably expired: re-run udemy_manifest.js, then build again)"
            if e.code in (401, 403, 404, 410):
                break
        except Exception as e:  # noqa: BLE001 - report and retry
            last = f"{type(e).__name__}: {e}"
        time.sleep(2 * (attempt + 1))
    try:
        tmp.unlink()
    except FileNotFoundError:
        pass
    return False, last


def fetch_bytes(url: str, timeout: int = 60) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Referer": "https://www.udemy.com/"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def filename_from_url(url: str) -> str:
    path = urllib.parse.urlparse(url).path
    return urllib.parse.unquote(path.rsplit("/", 1)[-1]) or "file"


# ----------------------------------------------------------------------------- HTML -> text
class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "section", "article", "header", "footer", "ul", "ol", "table", "tr",
             "blockquote", "pre", "figure", "figcaption", "dl", "dt", "dd", "hr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.href: list[str | None] = []
        self.link_text: list[str] = []
        self.pre = 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag == "br":
            self.out.append("\n")
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.out.append("\n\n" + "#" * min(int(tag[1]) + 1, 6) + " ")
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in ("td", "th"):
            self.out.append(" | ")
        elif tag == "img":
            alt = (a.get("alt") or "").strip()
            self.out.append(f" [image{': ' + alt if alt else ''}] ")
        elif tag == "a":
            self.href.append(a.get("href"))
            self.link_text.append("")
        elif tag == "pre":
            self.pre += 1
            self.out.append("\n")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.skip = max(0, self.skip - 1)
        elif tag == "a" and self.href:
            href = self.href.pop()
            text = self.link_text.pop().strip() if self.link_text else ""
            if href and href.startswith("http") and href not in text:
                self.out.append(f" ({href})")
        elif tag == "pre":
            self.pre = max(0, self.pre - 1)
            self.out.append("\n")
        elif tag in self.BLOCK or tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.out.append("\n")

    def handle_data(self, data):
        if self.skip:
            return
        if not self.pre:
            data = re.sub(r"\s+", " ", data)
        if self.link_text:
            self.link_text[-1] += data
        self.out.append(data)


def html_to_text(s: str) -> str:
    if not s:
        return ""
    if "<" not in s:
        return html.unescape(s).strip()
    p = _HTMLText()
    try:
        p.feed(s)
        p.close()
    except Exception:  # noqa: BLE001 - fall back to tag stripping
        return re.sub(r"<[^>]+>", " ", html.unescape(s)).strip()
    text = "".join(p.out)
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ----------------------------------------------------------------------------- captions
_TS = re.compile(r"(?:(\d+):)?(\d{1,2}):(\d{2})(?:[.,](\d{1,3}))?")


def parse_ts(s: str) -> float:
    m = _TS.search(s or "")
    if not m:
        return 0.0
    h, mi, se, ms = m.groups()
    return int(h or 0) * 3600 + int(mi) * 60 + int(se) + int((ms or "0").ljust(3, "0")) / 1000.0


def parse_vtt(text: str) -> list[dict]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    cues: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if "-->" in line:
            a, b = line.split("-->", 1)
            start, end = parse_ts(a), parse_ts(b.strip().split(" ")[0])
            i += 1
            buf = []
            while i < len(lines) and lines[i].strip():
                buf.append(lines[i])
                i += 1
            t = html.unescape(re.sub(r"<[^>]+>", "", " ".join(buf)))
            t = re.sub(r"\s+", " ", t).strip()
            if not t:
                continue
            if cues and t == cues[-1]["text"]:
                cues[-1]["end"] = max(cues[-1]["end"], end)
                continue
            if cues and t.startswith(cues[-1]["text"] + " "):  # rolling captions
                t = t[len(cues[-1]["text"]) + 1:]
            cues.append({"start": round(start, 2), "end": round(end, 2), "text": t})
        else:
            i += 1
    return cues


def paragraphs(segments: list[dict], target: float = 45.0) -> list[tuple[float, str]]:
    out, cur, start = [], [], None
    for s in segments:
        if start is None:
            start = float(s.get("start", 0))
        cur.append(s["text"].strip())
        span = float(s.get("end", 0)) - start
        sentence_end = re.search(r"[.!?][\"')\]]?$", s["text"].strip()) is not None
        if (span >= target and sentence_end) or span >= target * 2:
            out.append((start, " ".join(cur)))
            cur, start = [], None
    if cur:
        out.append((start or 0.0, " ".join(cur)))
    return out


def caption_locale(c: dict) -> str:
    loc = c.get("locale_id") or c.get("locale") or ""
    if isinstance(loc, dict):
        loc = loc.get("locale") or ""
    return str(loc)


def caption_is_auto(c: dict) -> bool:
    label = f"{c.get('source', '')} {c.get('title', '')} {c.get('video_label', '')}".lower()
    return "auto" in label


def pick_caption(captions, lang: str = "en") -> dict | None:
    caps = [c for c in (captions or []) if isinstance(c, dict) and c.get("url")]
    if not caps:
        return None
    for want_auto in (False, True):
        for c in caps:
            if caption_locale(c).lower().startswith(lang.lower()) and caption_is_auto(c) == want_auto:
                return c
    return caps[0]


# ----------------------------------------------------------------------------- OCR
def order_lines(items, min_score: float = 0.5) -> str:
    rows = []
    for box, text, score in items:
        try:
            score = float(score)
        except (TypeError, ValueError):
            score = 1.0
        if not text or score < min_score:
            continue
        xs = [float(p[0]) for p in box]
        ys = [float(p[1]) for p in box]
        rows.append((min(ys), max(ys), min(xs), str(text)))
    if not rows:
        return ""
    hs = sorted(r[1] - r[0] for r in rows)
    h = max(1.0, hs[len(hs) // 2])
    rows.sort(key=lambda r: ((r[0] + r[1]) / 2, r[2]))
    lines, cur, cy_avg = [], [], 0.0
    for r in rows:
        cy = (r[0] + r[1]) / 2
        if cur and abs(cy - cy_avg) > h * 0.6:
            lines.append(cur)
            cur = []
        cur.append(r)
        cy_avg = sum((x[0] + x[1]) / 2 for x in cur) / len(cur)
    if cur:
        lines.append(cur)
    return "\n".join(" ".join(x[3] for x in sorted(line, key=lambda x: x[2])) for line in lines)


class OCR:
    """Thin wrapper: RapidOCR (pip, CPU) first, then Tesseract; name is None if neither works."""

    def __init__(self, engine: str = "auto"):
        self.name = None
        self.error = None
        if engine == "none":
            return
        order = ["rapidocr", "tesseract"] if engine == "auto" else [engine]
        for name in order:
            try:
                getattr(self, "_init_" + name)()
                self.name = name
                break
            except Exception as e:  # noqa: BLE001
                self.error = f"{name}: {type(e).__name__}: {e}"

    def _init_rapidocr(self):
        try:
            from rapidocr_onnxruntime import RapidOCR  # bundled models
            self._new = False
        except ImportError:
            from rapidocr import RapidOCR  # newer package name
            self._new = True
        self._rapid = RapidOCR()

    def _init_tesseract(self):
        import pytesseract
        for p in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                  r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if os.path.exists(p):
                pytesseract.pytesseract.tesseract_cmd = p
        pytesseract.get_tesseract_version()
        self._tess = pytesseract

    def image_text(self, path) -> str:
        if not self.name:
            return ""
        path = str(path)
        if self.name == "rapidocr":
            if self._new:
                res = self._rapid(path)
                boxes = getattr(res, "boxes", None)
                txts = getattr(res, "txts", None)
                scores = getattr(res, "scores", None)
                if boxes is None or txts is None:
                    return ""
                items = list(zip(boxes, txts, scores if scores is not None else [1.0] * len(txts)))
            else:
                result, _ = self._rapid(path)
                items = [(b, t, s) for b, t, s in (result or [])]
            return order_lines(items)
        from PIL import Image
        with Image.open(path) as im:
            return self._tess.image_to_string(im.convert("RGB")).strip()

    def pil_text(self, img) -> str:
        if not self.name:
            return ""
        fd, tmp = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        try:
            img.convert("RGB").save(tmp)
            return self.image_text(tmp)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass


# ----------------------------------------------------------------------------- documents
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def docx_text(path) -> str:
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    paras = []
    for p in root.iter(W_NS + "p"):
        parts = []
        for node in p.iter():
            if node.tag == W_NS + "t" and node.text:
                parts.append(node.text)
            elif node.tag == W_NS + "tab":
                parts.append("\t")
            elif node.tag in (W_NS + "br", W_NS + "cr"):
                parts.append("\n")
        paras.append("".join(parts))
    return "\n".join(paras).strip()


def pptx_text(path) -> str:
    def num(n):
        m = re.search(r"(\d+)\.xml$", n)
        return int(m.group(1)) if m else 0

    def text_of(xml: bytes) -> str:
        root = ET.fromstring(xml)
        paras = ["".join(t.text or "" for t in p.iter(A_NS + "t")) for p in root.iter(A_NS + "p")]
        return "\n".join(x for x in paras if x.strip())

    out = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for s in sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=num):
            out.append(f"--- slide {num(s)} ---\n{text_of(z.read(s))}")
        for s in sorted((n for n in names if re.match(r"ppt/notesSlides/notesSlide\d+\.xml$", n)), key=num):
            t = text_of(z.read(s))
            if t.strip():
                out.append(f"--- speaker notes {num(s)} ---\n{t}")
    return "\n".join(out).strip()


def xlsx_text(path, max_rows: int = 3000, max_cols: int = 60) -> str:
    import openpyxl
    wb_f = openpyxl.load_workbook(path, data_only=False)
    try:
        wb_v = openpyxl.load_workbook(path, data_only=True)
    except Exception:  # noqa: BLE001
        wb_v = None
    out = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title] if wb_v is not None and ws.title in wb_v.sheetnames else None
        out.append(f"=== sheet: {ws.title} (used range {ws.dimensions}) ===")
        formulas = []
        rows, cols = min(ws.max_row or 0, max_rows), min(ws.max_column or 0, max_cols)
        for r in range(1, rows + 1):
            cells = []
            for c in range(1, cols + 1):
                cell = ws.cell(r, c)
                val = cell.value
                if type(val).__name__ == "ArrayFormula":
                    val = getattr(val, "text", str(val))
                if isinstance(val, str) and val.startswith("="):
                    cached = wv.cell(r, c).value if wv is not None else None
                    formulas.append(f"{cell.coordinate}: {val}  -> {cached!r}")
                    val = cached if cached is not None else val
                cells.append("" if val is None else str(val))
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                out.append(f"{r}\t" + "\t".join(cells))
        if (ws.max_row or 0) > max_rows or (ws.max_column or 0) > max_cols:
            out.append(f"... truncated: sheet is {ws.max_row} rows x {ws.max_column} columns")
        if formulas:
            out.append("--- formulas (cell: formula -> cached value) ---")
            out.extend(formulas[:5000])
    return "\n".join(out)


def xls_text(path) -> tuple[str, str]:
    try:
        import xlrd
    except ImportError:
        return "", "xlrd missing (pip install xlrd)"
    wb = xlrd.open_workbook(str(path))
    out = []
    for sh in wb.sheets():
        out.append(f"=== sheet: {sh.name} ({sh.nrows}x{sh.ncols}) ===")
        for r in range(min(sh.nrows, 3000)):
            vals = [str(v) for v in sh.row_values(r)[:60]]
            while vals and vals[-1] == "":
                vals.pop()
            if vals:
                out.append(f"{r + 1}\t" + "\t".join(vals))
    return "\n".join(out), "xlrd (values only)"


def pdf_text(path, ocr: OCR | None, max_pages: int = 400) -> tuple[str, str]:
    """Text layer via pypdf; pages with (almost) no text are OCR'd when an engine and pypdfium2 exist."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "", "pypdf missing (pip install pypdf)"
    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages[:max_pages]:
        try:
            texts.append((page.extract_text() or "").strip())
        except Exception as e:  # noqa: BLE001
            texts.append(f"[text extraction error: {e}]")
    empty = [i for i, t in enumerate(texts) if len(re.sub(r"\s", "", t)) < 10]
    method = "pypdf"
    if empty and ocr and ocr.name:
        try:
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(str(path))
            for i in empty:
                texts[i] = "(OCR) " + ocr.pil_text(doc[i].render(scale=2).to_pil())
            method = f"pypdf + {ocr.name} OCR for {len(empty)} image-only page(s)"
        except ImportError:
            method = f"pypdf ({len(empty)} image-only page(s) not OCR'd: pip install pypdfium2)"
    elif empty:
        method = f"pypdf ({len(empty)} image-only page(s) not OCR'd: no OCR engine)"
    return "\n".join(f"--- page {i + 1} ---\n{t}" for i, t in enumerate(texts)), method


def safe_unzip(zpath: Path, dest: Path, max_files: int = 3000, max_total: int = 2_000_000_000) -> list[Path]:
    out, total = [], 0
    with zipfile.ZipFile(zpath) as z:
        for info in z.infolist()[:max_files]:
            if info.is_dir():
                continue
            name = info.filename.replace("\\", "/")
            if name.startswith("/") or ".." in name.split("/") or re.match(r"^[A-Za-z]:", name):
                continue
            if "__MACOSX/" in name or name.split("/")[-1].startswith("._"):
                continue
            total += info.file_size
            if total > max_total:
                break
            target = dest.joinpath(*[safe_name(p, 50) for p in name.split("/") if p])
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                with z.open(info) as src, open(target, "wb") as dst:
                    shutil.copyfileobj(src, dst)
                out.append(target)
            except (OSError, zipfile.BadZipFile, RuntimeError) as e:  # long path, encrypted member, ...
                log(f"    skipped zip member {name}: {type(e).__name__}: {e}")
    return out


def extract_file(path: Path, ocr: OCR | None, base: Path, include_media: bool = False) -> list[dict]:
    """Extract text from one file (recursing into zips). Writes <file>.txt next to it."""
    ext = path.suffix.lower()
    if ext == ".zip":
        dest = path.with_name(path.stem + "_unzipped")
        results = [{"file": rel(path, base), "method": "zip", "text_file": None, "chars": 0}]
        try:
            for member in safe_unzip(path, dest):
                if member.suffix.lower() in MEDIA_EXT and not include_media:
                    continue
                results.extend(extract_file(member, ocr, base, include_media))
        except (zipfile.BadZipFile, OSError, RuntimeError) as e:
            results[0]["error"] = f"zip: {type(e).__name__}: {e}"
        return results
    text, method, err = "", "", None
    try:
        if ext == ".pdf":
            text, method = pdf_text(path, ocr)
        elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
            text, method = xlsx_text(path), "openpyxl (values, then formulas)"
        elif ext == ".xls":
            text, method = xls_text(path)
        elif ext == ".docx":
            text, method = docx_text(path), "docx xml"
        elif ext == ".pptx":
            text, method = pptx_text(path), "pptx xml"
        elif ext in (".html", ".htm"):
            text, method = html_to_text(read_text(path)), "html"
        elif ext in IMAGE_EXT:
            text, method = (ocr.image_text(path), f"{ocr.name} OCR") if ocr and ocr.name else ("", "image (no OCR engine)")
        elif ext in TEXT_EXT or ext == "":
            text, method = read_text(path, limit=5_000_000), "text"
        else:
            method = "not extracted (unsupported type)"
    except Exception as e:  # noqa: BLE001
        err = f"{type(e).__name__}: {e}"
        method = method or "error"
    rec = {"file": rel(path, base), "method": method, "text_file": None, "chars": len(text or "")}
    if err:
        rec["error"] = err
    if text and text.strip():
        tf = path.with_name(path.name + ".txt")
        tf.write_text(text, encoding="utf-8")
        rec["text_file"] = rel(tf, base)
    return [rec]


# ----------------------------------------------------------------------------- state
class State:
    def __init__(self, course_dir):
        self.d = Path(course_dir)
        self.sd = self.d / "state"

    @property
    def course(self) -> dict:
        p = self.sd / "course.json"
        return read_json(p) if p.exists() else {}

    @property
    def order(self) -> list[dict]:
        p = self.sd / "order.json"
        return read_json(p) if p.exists() else []

    def path(self, key: str) -> Path:
        return self.sd / "items" / f"{key}.json"

    def load(self, key: str) -> dict:
        p = self.path(key)
        return read_json(p) if p.exists() else {}

    def save(self, key: str, obj: dict) -> None:
        write_json(self.path(key), obj)

    def lecture_keys(self) -> list[str]:
        return [e["key"] for e in self.order if e.get("kind") == "lecture"]

    def key_for(self, lecture_id) -> str | None:
        lecture_id = str(lecture_id)
        for e in self.order:
            if e.get("kind") in ("lecture", "quiz") and str(e.get("id")) == lecture_id:
                return e["key"]
        return None


def require_state(course_dir) -> State:
    st = State(course_dir)
    if not (st.sd / "order.json").exists():
        sys.exit(f"{course_dir} has no state/order.json - run 'build' first")
    return st


# ----------------------------------------------------------------------------- build
def course_slug(m: dict) -> str:
    c = m.get("course") or {}
    if c.get("published_title"):
        return c["published_title"]
    mt = re.search(r"/course/([^/]+)", str(c.get("url") or ""))
    return mt.group(1) if mt else str(m.get("course_id") or "course")


def assets_from_download_urls(du) -> list[dict]:
    files = []
    if isinstance(du, dict):
        for kind, lst in du.items():
            for x in lst or []:
                if isinstance(x, dict) and x.get("file"):
                    files.append({"kind": kind, "label": x.get("label"), "url": x["file"]})
    return files


def unique_name(name: str, used: set) -> str:
    """Deterministic de-duplication of file names inside one lecture folder."""
    cand, i = name, 2
    while cand.lower() in used:
        stem, ext = os.path.splitext(name)
        cand = f"{stem}_{i}{ext}"
        i += 1
    used.add(cand.lower())
    return cand


def fetch_and_extract(url: str, folder: Path, name: str, ocr, base: Path, include_media: bool,
                      used: set | None = None) -> list[dict]:
    name = safe_name(name or filename_from_url(url), 80)
    if used is not None:
        name = unique_name(name, used)
    if Path(name).suffix.lower() in MEDIA_EXT and not include_media:
        return [{"file": None, "name": name, "method": "skipped media file (use --include-media)", "chars": 0}]
    dest = folder / name
    ok, err = download(url, dest)
    if not ok:
        return [{"file": None, "name": name, "error": err, "chars": 0}]
    return extract_file(dest, ocr, base, include_media)


def build_lecture(it: dict, key: str, prev: dict, d: Path, origin: str, slug: str, a, ocr) -> dict:
    det = it.get("details") or {}
    asset = det.get("asset") or it.get("asset") or {}
    atype = asset.get("asset_type") or "?"
    lec = dict(prev)
    lec.update({
        "kind": "lecture", "key": key, "id": it.get("id"), "n": it.get("lecture_n"),
        "title": it.get("title") or det.get("title") or "", "section_n": it.get("section_n"),
        "section_title": it.get("section_title"), "asset_type": atype,
        "duration_s": asset.get("length") or asset.get("time_estimation"),
        "url": f"{origin}/course/{slug}/learn/lecture/{it.get('id')}",
        "description": html_to_text(det.get("description") or it.get("description") or ""),
        "thumbnail_sprite": asset.get("thumbnail_sprite"), "drm": asset.get("course_is_drmed"),
        "errors": [],
    })
    if atype == "Video":
        lec["video_download_urls"] = (asset.get("download_urls") or {}).get("Video") or []
    # transcript from Udemy captions (a Whisper transcript is kept unless --refresh)
    if atype in ("Video", "VideoMashup"):
        have_whisper = str((prev.get("transcript") or {}).get("source", "")).startswith("whisper")
        cap = pick_caption(asset.get("captions"), a.lang)
        if cap and (a.refresh or not have_whisper):
            loc = caption_locale(cap) or "xx"
            vtt = d / "captions" / f"{key}.{safe_name(loc, 12)}.vtt"
            if a.refresh and vtt.exists():
                vtt.unlink()
            ok, err = download(cap["url"], vtt)
            if ok:
                lec["transcript"] = {"source": f"udemy captions {loc} {'auto' if caption_is_auto(cap) else 'manual'}",
                                     "segments": parse_vtt(read_text(vtt))}
            else:
                lec["errors"].append(f"captions: {err}")
        lec["needs_transcription"] = not (lec.get("transcript") or {}).get("segments")
    if atype == "Article" or asset.get("body"):
        lec["article"] = html_to_text(asset.get("body") or "")
    # presentation slides (images)
    slide_urls = list(asset.get("slide_urls") or [])
    for s in asset.get("slides") or []:
        if isinstance(s, dict) and s.get("url") and s["url"] not in slide_urls:
            slide_urls.append(s["url"])
    if slide_urls:
        slides = []
        for i, u in enumerate(slide_urls, 1):
            ext = Path(filename_from_url(u)).suffix.lower() or ".jpg"
            dest = d / "slides" / key / f"{i:03d}{ext if ext in IMAGE_EXT else '.jpg'}"
            ok, err = download(u, dest)
            if not ok:
                lec["errors"].append(f"slide {i}: {err}")
                continue
            old = next((x for x in prev.get("slides") or [] if x.get("file") == rel(dest, d)), None)
            text = old["text"] if old and old.get("text") else (ocr.image_text(dest) if ocr and ocr.name else "")
            slides.append({"n": i, "file": rel(dest, d), "text": text})
        lec["slides"] = slides
    # non-video lecture files (PDF e-books, presentations, files)
    files, used = [], set()
    if atype != "Video":
        for f in assets_from_download_urls(asset.get("download_urls")):
            files += fetch_and_extract(f["url"], d / "resources" / key, asset.get("filename") or asset.get("title"),
                                       ocr, d, a.include_media, used)
    lec["asset_files"] = files
    links = [asset["external_url"]] if asset.get("external_url") else []
    # supplementary resources
    resources = []
    for r in it.get("resources") or it.get("supplementary_assets") or []:
        rec = {"title": r.get("title") or r.get("filename"), "type": r.get("asset_type"), "files": []}
        if r.get("error"):
            rec["error"] = r["error"]
        if r.get("external_url"):
            links.append(r["external_url"])
            rec["url"] = r["external_url"]
        for f in assets_from_download_urls(r.get("download_urls")):
            rec["files"] += fetch_and_extract(f["url"], d / "resources" / key, r.get("filename") or r.get("title"),
                                              ocr, d, a.include_media, used)
        if r.get("body"):
            rec["text"] = html_to_text(r["body"])
        resources.append(rec)
    lec["resources"] = resources
    lec["external_links"] = links
    return lec


def build_quiz(it: dict, key: str, prev: dict) -> dict:
    qs = []
    for x in it.get("assessments") or []:
        p = x.get("prompt") or {}
        qs.append({
            "type": x.get("assessment_type"),
            "question": html_to_text(p.get("question") or x.get("question_plain") or ""),
            "answers": [html_to_text(v) for v in p.get("answers") or []],
            "feedbacks": [html_to_text(v) for v in p.get("feedbacks") or []],
            "correct": x.get("correct_response") or [],
            "explanation": html_to_text(p.get("explanation") or ""),
            "section": x.get("section"),
        })
    q = dict(prev)
    q.update({"kind": "quiz", "key": key, "id": it.get("id"), "n": it.get("quiz_n"), "title": it.get("title") or "",
              "quiz_type": it.get("type"), "section_n": it.get("section_n"), "section_title": it.get("section_title"),
              "description": html_to_text(it.get("description") or ""), "questions": qs})
    return q


def cmd_build(a) -> None:
    m = read_json(a.manifest)
    if not isinstance(m, dict) or "curriculum" not in m:
        sys.exit("not a udemy_manifest.js manifest (no 'curriculum')")
    d = Path(a.out)
    d.mkdir(parents=True, exist_ok=True)
    st = State(d)
    src = Path(a.manifest).resolve()
    if src != (d / "manifest.json").resolve():
        shutil.copy2(src, d / "manifest.json")
    ocr = OCR(a.ocr)
    log(f"OCR engine: {ocr.name or 'none'}" + (f" ({ocr.error})" if not ocr.name and ocr.error else ""))
    course = m.get("course") or {}
    origin = m.get("origin") or "https://www.udemy.com"
    slug = course_slug(m)
    write_json(st.sd / "course.json", {
        "id": m.get("course_id") or course.get("id"), "title": course.get("title"), "slug": slug,
        "url": f"{origin}/course/{slug}/", "headline": course.get("headline"),
        "instructors": [i.get("title") or i.get("display_name") for i in course.get("visible_instructors") or []],
        "description": html_to_text(course.get("description") or ""), "content_info": course.get("content_info"),
        "last_update_date": course.get("last_update_date"), "manifest_generated_at": m.get("generated_at"),
        "manifest_errors": m.get("errors") or [], "built_at": dt.datetime.now().isoformat(timespec="seconds"),
        "tool": f"udemy_process.py {VERSION}"})
    order, sec_n, sec_title, lec_n, quiz_n = [], 0, "", 0, 0
    items = m["curriculum"]
    n_items = sum(1 for x in items if x.get("_class") in ("lecture", "quiz"))
    done = 0
    for it in items:
        cls = it.get("_class")
        if cls == "chapter":
            sec_n += 1
            sec_title = it.get("title") or ""
            order.append({"kind": "section", "n": sec_n, "title": sec_title})
            continue
        it.setdefault("section_n", sec_n)
        it.setdefault("section_title", sec_title)
        if cls == "lecture":
            lec_n += 1
            it.setdefault("lecture_n", lec_n)
            key = f"L{int(it['lecture_n']):03d}_{it.get('id')}"
            obj = build_lecture(it, key, st.load(key), d, origin, slug, a, ocr)
            order.append({"kind": "lecture", "key": key, "id": it.get("id")})
        elif cls == "quiz":
            quiz_n += 1
            it.setdefault("quiz_n", quiz_n)
            key = f"Q{int(it['quiz_n']):02d}_{it.get('id')}"
            obj = build_quiz(it, key, st.load(key))
            order.append({"kind": "quiz", "key": key, "id": it.get("id")})
        else:
            key = f"X_{it.get('id')}"
            obj = {"kind": "other", "key": key, "id": it.get("id"), "title": it.get("title"), "class": cls,
                   "section_n": it.get("section_n"), "section_title": it.get("section_title")}
            order.append({"kind": "other", "key": key, "id": it.get("id")})
        st.save(key, obj)
        if cls in ("lecture", "quiz"):
            done += 1
            if done % 10 == 0 or done == n_items:
                log(f"  {done}/{n_items} items processed")
    write_json(st.sd / "order.json", order)
    assemble(d)
    cmd_status(argparse.Namespace(course_dir=str(d)))


# ----------------------------------------------------------------------------- frames
def gray_small(img, size=(160, 90)) -> bytes:
    return img.convert("L").resize(size).tobytes()


def diff_frac(a, b, thr: int = 24) -> float:
    return sum(1 for x, y in zip(a, b) if abs(x - y) > thr) / max(1, len(a))


class SceneDetector:
    """Same rule as udemy_frames.js. Motion is judged per 8x8 block of the 160x90 thumbnail:
    a block is dynamic if it moves now, moved when the reference frame was kept, or moved in
    2+ of the last 6 pairs. A frame is kept when >= lo of the static pixels changed."""

    def __init__(self, lo=0.015, motion=0.02, motion_gap=45.0, max_gap=120.0, hist=6, thr=24,
                 size=(160, 90), block=8):
        self.lo, self.motion, self.motion_gap, self.max_gap, self.k, self.thr = lo, motion, motion_gap, max_gap, hist, thr
        self.w, self.h = size
        self.bw = (self.w + block - 1) // block
        self.block_of = [(i // self.w // block) * self.bw + (i % self.w) // block for i in range(self.w * self.h)]
        self.nb = self.bw * ((self.h + block - 1) // block)
        self.hist: list[list[int]] = []
        self.last = None
        self.last_t = -1e9
        self.mvb = None        # moving blocks of the latest decide() call
        self.last_mvb = None   # moving blocks when the reference frame was kept

    def decide(self, t: float, a, b=None) -> tuple[bool, dict]:
        mvb = None
        if b is not None:
            mvb = [0] * self.nb
            for i, (x, y) in enumerate(zip(a, b)):
                if abs(x - y) > self.thr:
                    mvb[self.block_of[i]] = 1
            self.hist = (self.hist + [mvb])[-self.k:]
        self.mvb = mvb
        if self.last is None:
            return True, {"static_change": 1.0, "motion": 0.0, "change": 1.0}
        dyn_block = [0] * self.nb
        for bi in range(self.nb):
            dyn_block[bi] = int((mvb is not None and mvb[bi] == 1) or (self.last_mvb is not None and self.last_mvb[bi] == 1)
                                or sum(h[bi] for h in self.hist) >= 2)
        stat = stat_ch = ch = dyn = 0
        for i, x in enumerate(a):
            changed = abs(x - self.last[i]) > self.thr
            ch += changed
            if dyn_block[self.block_of[i]]:
                dyn += 1
            else:
                stat += 1
                stat_ch += changed
        n = max(1, len(a))
        st = {"static_change": stat_ch / max(1, stat), "motion": dyn / n, "change": ch / n}
        keep = (st["static_change"] >= self.lo
                or (st["motion"] >= self.motion and t - self.last_t >= self.motion_gap and st["change"] >= self.lo)
                or (t - self.last_t >= self.max_gap and st["change"] >= 0.005))
        return keep, st

    def kept(self, t: float, thumb) -> None:
        self.last, self.last_t, self.last_mvb = thumb, t, self.mvb


def parse_sprite_vtt(text: str, base_url: str) -> list[tuple[float, float, str, tuple | None]]:
    out = []
    lines = text.replace("\r\n", "\n").split("\n")
    for i, line in enumerate(lines):
        if "-->" not in line:
            continue
        a, b = line.split("-->", 1)
        payload = next((x.strip() for x in lines[i + 1:i + 3] if x.strip()), "")
        if not payload:
            continue
        url, _, frag = payload.partition("#")
        box = None
        mt = re.search(r"xywh=(\d+),(\d+),(\d+),(\d+)", frag)
        if mt:
            x, y, w, h = map(int, mt.groups())
            box = (x, y, x + w, y + h)
        out.append((parse_ts(a), parse_ts(b.strip().split(" ")[0]), urllib.parse.urljoin(base_url, url), box))
    return out


def sprite_urls(sp) -> tuple[str | None, str | None]:
    vtt = img = None
    if isinstance(sp, dict):
        for v in sp.values():
            if isinstance(v, str) and re.search(r"\.vtt(\?|$)", v):
                vtt = vtt or v
            elif isinstance(v, str) and re.search(r"\.(jpe?g|png|webp)(\?|$)", v):
                img = img or v
    elif isinstance(sp, str) and ".vtt" in sp:
        vtt = sp
    return vtt, img


def cmd_sprite_times(a) -> None:
    from PIL import Image
    st = require_state(a.course_dir)
    out_path = Path(a.course_dir) / "frame_times.json"
    result = read_json(out_path) if out_path.exists() else {}
    for key in st.lecture_keys():
        lec = st.load(key)
        if lec.get("asset_type") != "Video" or (str(lec.get("id")) in result and not a.refresh):
            continue
        vtt_url, _ = sprite_urls(lec.get("thumbnail_sprite"))
        if not vtt_url:
            continue
        try:
            cues = parse_sprite_vtt(fetch_bytes(vtt_url).decode("utf-8", "replace"), vtt_url)
            sheets: dict[str, object] = {}
            times, last, last_t = [], None, -1e9
            for s, e, url, box in cues:
                if url not in sheets:
                    sheets[url] = Image.open(io.BytesIO(fetch_bytes(url))).convert("L")
                tile = sheets[url].crop(box) if box else sheets[url]
                g = gray_small(tile)
                dfr = 1.0 if last is None else diff_frac(g, last)
                keep = last is None or dfr >= a.lo or s - last_t >= a.max_gap
                if keep:
                    times.append(round(min(max(s, e - 0.2), s + max(0.5, (e - s) / 2)), 1))
                    last, last_t = g, s
            result[str(lec["id"])] = times
            log(f"{key}: {len(times)} timestamps from {len(cues)} sprite tiles")
        except Exception as ex:  # noqa: BLE001
            log(f"{key}: sprite analysis failed ({type(ex).__name__}: {ex}); interval capture will be used")
        write_json(out_path, result)
    log(f"wrote {out_path}")


def cmd_frame_plan(a) -> None:
    st = require_state(a.course_dir)
    tp = Path(a.course_dir) / "frame_times.json"
    times = read_json(tp) if tp.exists() else {}
    pending = []
    for key in st.lecture_keys():
        lec = st.load(key)
        if lec.get("asset_type") != "Video" or lec.get("frames") is not None or lec.get("frames_error"):
            continue
        t = times.get(str(lec.get("id")))
        call = (f"udfStart({{times:{json.dumps(t, separators=(',', ':'))}}})" if t
                else f"udfStart({{interval:{a.interval}}})")
        pending.append(f"{key}\t{lec.get('url')}\t{call}")
    log(f"{len(pending)} video lectures still need frames")
    for line in pending[: a.limit]:
        print(line)


def cmd_frames(a) -> None:
    st = require_state(a.course_dir)
    d = Path(a.course_dir)
    ocr = OCR(a.ocr)
    log(f"OCR engine: {ocr.name or 'none'}" + (f" ({ocr.error})" if not ocr.name and ocr.error else ""))
    paths = []
    for pat in a.frames_json:
        paths += sorted(glob.glob(pat)) or [pat]
    for fj in paths:
        try:
            data = read_json(fj)
        except Exception as ex:  # noqa: BLE001
            log(f"skip {fj}: {ex}")
            continue
        key = st.key_for(data.get("lecture_id"))
        if not key:
            log(f"skip {fj}: lecture {data.get('lecture_id')} is not in this course")
            continue
        lec = st.load(key)
        outdir = d / "frames" / key
        outdir.mkdir(parents=True, exist_ok=True)
        frames, prev_i = [], None
        for fr in data.get("frames") or []:
            t = float(fr.get("t", 0))
            raw = str(fr.get("data", ""))
            if "," not in raw:
                continue
            fn = outdir / f"t{int(round(t * 10)):06d}.jpg"
            fn.write_bytes(base64.b64decode(raw.split(",", 1)[1]))
            text = ocr.image_text(fn) if ocr.name else ""
            rec = {"t": t, "file": rel(fn, d), "text": text, "dup_text": False, "diff": fr.get("diff")}
            if text.strip() and prev_i is not None:
                prev = frames[prev_i]["text"]
                if similar(text, prev) >= 0.9:
                    rec["dup_text"] = True          # same slide again
                elif " ".join(prev.split()) in " ".join(text.split()):
                    frames[prev_i]["dup_text"] = True  # previous was an earlier build of this slide
            frames.append(rec)
            if text.strip() and not rec["dup_text"]:
                prev_i = len(frames) - 1
        if a.append and lec.get("frames"):
            have = {round(float(f["t"]), 1) for f in frames}
            frames = sorted([f for f in lec["frames"] if round(float(f["t"]), 1) not in have] + frames,
                            key=lambda f: float(f["t"]))
        lec["frames"] = frames
        lec.pop("frames_error", None)
        lec["frames_info"] = {"source": data.get("tool"), "checked": data.get("checked"),
                              "median_source_width": data.get("median_source_width"), "ocr": ocr.name}
        st.save(key, lec)
        raw_dir = d / "frames" / "_raw"
        raw_dir.mkdir(parents=True, exist_ok=True)
        if Path(fj).resolve().parent != raw_dir.resolve():
            dest = raw_dir / Path(fj).name
            if dest.exists():
                dest = dest.with_name(f"{dest.stem}_{int(time.time())}{dest.suffix}")
            shutil.move(fj, dest)
        log(f"{key}: {len(frames)} frames stored, {sum(1 for f in frames if f['text'] and not f['dup_text'])} with new text")
    assemble(d)


def cmd_skip_frames(a) -> None:
    st = require_state(a.course_dir)
    key = st.key_for(a.lecture)
    if not key:
        sys.exit(f"lecture {a.lecture} not found")
    lec = st.load(key)
    lec["frames_error"] = a.reason
    st.save(key, lec)
    log(f"{key}: frames marked as not capturable ({a.reason})")


def cmd_screen_capture(a) -> None:
    from PIL import ImageGrab
    st = require_state(a.course_dir)
    if not st.key_for(a.lecture):
        sys.exit(f"lecture {a.lecture} not found")
    bbox = tuple(int(x) for x in a.bbox.split(",")) if a.bbox else None
    det = SceneDetector(a.lo, a.motion, a.motion_gap, a.max_gap)
    frames, prev = [], None
    t0 = time.time()
    log(f"capturing for {a.seconds}s every {a.every}s (video speed x{a.speed}) ...")
    while time.time() - t0 < a.seconds:
        tick = time.time()
        img = ImageGrab.grab(bbox=bbox, all_screens=a.all_screens)
        vt = round(a.start_at + (tick - t0) * a.speed, 1)
        g = gray_small(img)
        keep, st = det.decide(vt, g, prev)
        if keep:
            buf = io.BytesIO()
            img.convert("RGB").save(buf, "JPEG", quality=85)
            frames.append({"t": vt, "diff": round(st["static_change"], 3), "motion": round(st["motion"], 3),
                           "w": img.width, "h": img.height,
                           "data": "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()})
            det.kept(vt, g)
        prev = g
        time.sleep(max(0.0, a.every - (time.time() - tick)))
    out = Path(a.course_dir) / "frames" / "_raw" / f"screen_{a.lecture}_{int(t0)}.json"
    write_json(out, {"tool": "udemy_process.py screen-capture", "lecture_id": str(a.lecture),
                     "checked": None, "median_source_width": frames[0]["w"] if frames else 0, "frames": frames})
    log(f"{len(frames)} frames -> {out}\nnext: python udemy_process.py frames \"{a.course_dir}\" \"{out}\"")


# ----------------------------------------------------------------------------- whisper fallback
def cmd_transcribe(a) -> None:
    st = require_state(a.course_dir)
    d = Path(a.course_dir)
    key = st.key_for(a.lecture)
    if not key:
        sys.exit(f"lecture {a.lecture} not found")
    lec = st.load(key)
    media = Path(a.media) if a.media else None
    if a.download:
        urls = lec.get("video_download_urls") or []
        if not urls:
            sys.exit("Udemy offers no download for this lecture; record the audio while it plays and pass --media")

        def res(u):
            digits = re.sub(r"\D", "", str(u.get("label") or ""))
            return int(digits) if digits else 10 ** 6
        best = sorted(urls, key=res)[0]
        media = d / "media" / f"{key}.mp4"
        ok, err = download(best["file"], media)
        if not ok:
            sys.exit(f"download failed: {err}")
    if not media or not media.exists():
        sys.exit("give --media FILE or --download")
    from faster_whisper import WhisperModel
    log(f"transcribing {media} with faster-whisper {a.model} (CPU int8) ...")
    model = WhisperModel(a.model, device="cpu", compute_type="int8")
    segs, info = model.transcribe(str(media), vad_filter=True, language=a.language)
    segments = [{"start": round(s.start, 2), "end": round(s.end, 2), "text": s.text.strip()} for s in segs]
    lec["transcript"] = {"source": f"whisper {a.model}", "segments": segments}
    lec["needs_transcription"] = False
    st.save(key, lec)
    if a.download and not a.keep_media:
        media.unlink(missing_ok=True)
    log(f"{key}: {len(segments)} segments")
    assemble(d)


# ----------------------------------------------------------------------------- assemble
def render_transcript(lec: dict) -> str:
    tr = lec.get("transcript") or {}
    events = [(s, 0, "p", t) for s, t in paragraphs(tr.get("segments") or [])]
    for f in lec.get("frames") or []:
        if (f.get("text") or "").strip() and not f.get("dup_text"):
            events.append((float(f["t"]), 1, "f", f))
    events.sort(key=lambda e: (e[0], e[1]))
    out = []
    for t, _, kind, x in events:
        if kind == "p":
            out.append(f"[{fmt_ts(t)}] {x}")
        else:
            body = "\n".join("> " + ln for ln in x["text"].splitlines() if ln.strip())
            out.append(f"> [{fmt_ts(t)}] on screen ({x['file']}):\n{body}")
    return "\n\n".join(out)


def render_files(files: list[dict], d: Path) -> list[str]:
    out = []
    for f in files:
        if f.get("error") or not f.get("file"):
            out.append(f"- {f.get('name') or f.get('file')}: {f.get('error') or f.get('method')}")
            continue
        out.append(f"- `{f['file']}` ({f.get('method')}, {f.get('chars', 0):,} chars"
                   + (f"; full text `{f['text_file']}`" if f.get("text_file") else "") + ")")
        if f.get("text_file"):
            txt = read_text(d / f["text_file"])
            cut = len(txt) > EMBED_LIMIT
            out.append("```text\n" + txt[:EMBED_LIMIT].replace("```", "'''") +
                       ("\n... [truncated, see full text file]" if cut else "") + "\n```")
    return out


def render_item(x: dict, d: Path) -> str:
    if x.get("kind") == "quiz":
        lines = [f"## {x['key'].split('_')[0]} - Quiz: {x.get('title')}",
                 f"Section {x.get('section_n')}: {x.get('section_title')} - {x.get('quiz_type') or 'quiz'}", ""]
        if x.get("description"):
            lines += [x["description"], ""]
        for i, q in enumerate(x.get("questions") or [], 1):
            lines.append(f"{i}. {q['question']}")
            letters = "abcdefghijklmnopqrstuvwxyz"
            correct = {str(c).lower() for c in q.get("correct") or []}
            for j, ans in enumerate(q.get("answers") or []):
                mark = " [correct]" if letters[j] in correct else ""
                fb = (q.get("feedbacks") or [""] * 99)[j] if j < len(q.get("feedbacks") or []) else ""
                lines.append(f"   {letters[j]}) {ans}{mark}" + (f"  - feedback: {fb}" if fb else ""))
            if q.get("explanation"):
                lines.append(f"   Explanation: {q['explanation']}")
            lines.append("")
        return "\n".join(lines) + "\n"
    if x.get("kind") == "other":
        return f"## {x.get('title')} ({x.get('class')}, not extracted)\n"
    dur = f" - {fmt_ts(x['duration_s'])}" if x.get("duration_s") else ""
    lines = [f"## {x['key'].split('_')[0]} - {x.get('title')}",
             f"Section {x.get('section_n')}: {x.get('section_title')} - {x.get('asset_type')}{dur} - {x.get('url')}", ""]
    if x.get("description"):
        lines += ["**Lecture notes**", x["description"], ""]
    tr = x.get("transcript") or {}
    if tr.get("segments") or x.get("frames"):
        src = tr.get("source") or "no transcript"
        lines += [f"**Transcript** ({src}; quoted blocks are on-screen text from OCR)", render_transcript(x), ""]
    elif x.get("asset_type") == "Video":
        lines += ["**Transcript**: none yet (no captions; see needs_transcription in status)", ""]
    if x.get("frames_error"):
        lines += [f"_Frames not captured: {x['frames_error']}_", ""]
    if x.get("article"):
        lines += ["**Article**", x["article"], ""]
    if x.get("slides"):
        lines += ["**Slides (OCR)**"]
        for s in x["slides"]:
            lines += [f"--- slide {s['n']} (`{s['file']}`) ---", s.get("text") or "(no text found)"]
        lines.append("")
    if x.get("asset_files"):
        lines += ["**Lecture file**"] + render_files(x["asset_files"], d) + [""]
    if x.get("resources"):
        lines.append("**Resources**")
        for r in x["resources"]:
            lines.append(f"- {r.get('title')} ({r.get('type')})" + (f" - {r['url']}" if r.get("url") else "")
                         + (f" - error: {r['error']}" if r.get("error") else ""))
            if r.get("text"):
                lines.append("  " + r["text"].replace("\n", "\n  "))
            lines += ["  " + ln for ln in "\n".join(render_files(r.get("files") or [], d)).split("\n") if ln]
        lines.append("")
    if x.get("external_links"):
        lines += ["**External links**"] + [f"- {u}" for u in x["external_links"]] + [""]
    if x.get("errors"):
        lines += ["**Extraction errors**"] + [f"- {e}" for e in x["errors"]] + [""]
    return "\n".join(lines) + "\n"


def assemble(course_dir) -> None:
    d = Path(course_dir)
    st = require_state(d)
    c = st.course
    (d / "lectures").mkdir(exist_ok=True)
    for old in (d / "lectures").glob("*.md"):
        old.unlink()
    toc, body, index, words = [], [], [], 0
    for e in st.order:
        if e["kind"] == "section":
            toc.append(f"\n**Section {e['n']}: {e['title']}**")
            body.append(f"\n# Section {e['n']}: {e['title']}\n")
            continue
        x = st.load(e["key"])
        md = render_item(x, d)
        fname = f"{e['key']}_{safe_name(x.get('title'), 40)}.md"
        (d / "lectures" / fname).write_text(md, encoding="utf-8")
        body.append(md)
        w = len(md.split())
        words += w
        flag = ""
        if x.get("kind") == "lecture" and x.get("asset_type") == "Video":
            flag = " (no transcript)" if x.get("needs_transcription") else ""
        toc.append(f"- {e['key'].split('_')[0]} {x.get('title')}{flag}")
        index.append({"key": e["key"], "kind": x.get("kind"), "title": x.get("title"), "section_n": x.get("section_n"),
                      "file": f"lectures/{fname}", "words": w})
    head = [f"# {c.get('title')}", "",
            f"Course: {c.get('url')}  ", f"Instructors: {', '.join(i for i in c.get('instructors') or [] if i)}  ",
            f"Built: {c.get('built_at')} with {c.get('tool')} (personal copy - do not share or publish)", ""]
    if c.get("headline"):
        head += [c["headline"], ""]
    if c.get("description"):
        head += ["## Course description", c["description"], ""]
    text = "\n".join(head) + "\n## Contents\n" + "\n".join(toc) + "\n\n" + "\n".join(body)
    (d / "course.md").write_text(text, encoding="utf-8")
    write_json(d / "course_index.json", {"course": c.get("title"), "words": words, "items": index})


def cmd_assemble(a) -> None:
    assemble(a.course_dir)
    log(f"assembled {Path(a.course_dir) / 'course.md'}")


def cmd_status(a) -> None:
    st = require_state(a.course_dir)
    d = Path(a.course_dir)
    lec = [st.load(k) for k in st.lecture_keys()]
    vids = [x for x in lec if x.get("asset_type") == "Video"]
    quizzes = [e for e in st.order if e["kind"] == "quiz"]
    res_files = [f for x in lec for r in x.get("resources") or [] for f in r.get("files") or []]
    res_files += [f for x in lec for f in x.get("asset_files") or []]
    errors = [f"{x['key']}: {err}" for x in lec for err in x.get("errors") or []]
    errors += [f"{x['key']}: {f.get('name') or f.get('file')}: {f['error']}" for x in lec
               for r in x.get("resources") or [] for f in r.get("files") or [] if f.get("error")]
    errors += [f"{x['key']}: resource '{r.get('title')}': {r['error']}" for x in lec
               for r in x.get("resources") or [] if r.get("error")]
    no_tr = [x["key"] for x in vids if x.get("needs_transcription")]
    frames_done = sum(1 for x in vids if x.get("frames") is not None)
    frames_skip = sum(1 for x in vids if x.get("frames_error"))
    rows = [
        ("course", st.course.get("title")),
        ("lectures", len(lec)),
        ("video lectures", len(vids)),
        ("videos with transcript", len(vids) - len(no_tr)),
        ("videos needing transcription", len(no_tr)),
        ("videos with frames OCR'd", frames_done),
        ("videos with frames skipped", frames_skip),
        ("articles", sum(1 for x in lec if x.get("article"))),
        ("slide decks", sum(1 for x in lec if x.get("slides"))),
        ("resource files extracted", sum(1 for f in res_files if f.get("text_file"))),
        ("resource files failed", sum(1 for f in res_files if f.get("error"))),
        ("quizzes", len(quizzes)),
        ("errors", len(errors)),
        ("course.md words", read_json(d / "course_index.json").get("words") if (d / "course_index.json").exists() else 0),
        ("strategies spec", "present" if (d / "strategies.md").exists() else "not yet"),
    ]
    for k, v in rows:
        print(f"{k:30s} {v}")
    if no_tr:
        print("needs transcription:", ", ".join(no_tr[:40]) + (" ..." if len(no_tr) > 40 else ""))
    for e in errors[:15]:
        print("error:", e)
    write_json(d / "state" / "needs_transcription.json", no_tr)


# ----------------------------------------------------------------------------- CLI
def main(argv=None) -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--version", action="version", version=VERSION)
    sub = p.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="build the course folder from a manifest")
    b.add_argument("manifest")
    b.add_argument("--out", required=True, help="course folder to create/update")
    b.add_argument("--lang", default="en", help="caption language prefix (default en)")
    b.add_argument("--ocr", default="auto", choices=["auto", "rapidocr", "tesseract", "none"])
    b.add_argument("--include-media", action="store_true", help="also download audio/video resource files")
    b.add_argument("--refresh", action="store_true", help="re-download captions even if present")
    b.set_defaults(func=cmd_build)

    s = sub.add_parser("sprite-times", help="slide-change timestamps from thumbnail sprites")
    s.add_argument("course_dir")
    s.add_argument("--lo", type=float, default=0.015, help="share of changed pixels that marks a new scene")
    s.add_argument("--max-gap", type=float, default=90.0, help="add a timestamp at least this often")
    s.add_argument("--refresh", action="store_true")
    s.set_defaults(func=cmd_sprite_times)

    fp = sub.add_parser("frame-plan", help="video lectures still lacking frames")
    fp.add_argument("course_dir")
    fp.add_argument("--interval", type=float, default=8.0)
    fp.add_argument("--limit", type=int, default=500)
    fp.set_defaults(func=cmd_frame_plan)

    f = sub.add_parser("frames", help="store + OCR frames JSON files from udemy_frames.js")
    f.add_argument("course_dir")
    f.add_argument("frames_json", nargs="+", help="files or glob patterns, e.g. %%USERPROFILE%%/Downloads/frames_*.json")
    f.add_argument("--ocr", default="auto", choices=["auto", "rapidocr", "tesseract", "none"])
    f.add_argument("--append", action="store_true", help="merge with frames already stored for the lecture")
    f.set_defaults(func=cmd_frames)

    sk = sub.add_parser("skip-frames", help="mark a lecture whose frames cannot be captured")
    sk.add_argument("course_dir")
    sk.add_argument("--lecture", required=True)
    sk.add_argument("--reason", required=True)
    sk.set_defaults(func=cmd_skip_frames)

    sc = sub.add_parser("screen-capture", help="fallback: screenshots while a lecture plays")
    sc.add_argument("course_dir")
    sc.add_argument("--lecture", required=True)
    sc.add_argument("--seconds", type=float, required=True, help="wall-clock seconds to capture")
    sc.add_argument("--every", type=float, default=1.0)
    sc.add_argument("--speed", type=float, default=1.0, help="player speed, to map to video time")
    sc.add_argument("--start-at", type=float, default=0.0, help="video time when capture starts")
    sc.add_argument("--bbox", help="x1,y1,x2,y2 screen region of the video")
    sc.add_argument("--all-screens", action="store_true")
    sc.add_argument("--lo", type=float, default=0.015)
    sc.add_argument("--motion", type=float, default=0.02)
    sc.add_argument("--motion-gap", type=float, default=45.0)
    sc.add_argument("--max-gap", type=float, default=120.0)
    sc.set_defaults(func=cmd_screen_capture)

    t = sub.add_parser("transcribe", help="Whisper transcript for a lecture without captions")
    t.add_argument("course_dir")
    t.add_argument("--lecture", required=True)
    t.add_argument("--media", help="audio/video file of the lecture")
    t.add_argument("--download", action="store_true", help="use Udemy's own download link if the course allows it")
    t.add_argument("--model", default="small.en")
    t.add_argument("--language", default="en")
    t.add_argument("--keep-media", action="store_true")
    t.set_defaults(func=cmd_transcribe)

    asm = sub.add_parser("assemble", help="rebuild course.md and lectures/*.md")
    asm.add_argument("course_dir")
    asm.set_defaults(func=cmd_assemble)

    stt = sub.add_parser("status", help="what is done / missing")
    stt.add_argument("course_dir")
    stt.set_defaults(func=cmd_status)

    a = p.parse_args(argv)
    a.func(a)


if __name__ == "__main__":
    main()
